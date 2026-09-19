from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from backend.main import app, DOCX_MIME

client = TestClient(app)


def test_docx_editable_a4_calibri_and_saved():
    response = client.post('/api/docx', json={'title': 'Investigación del Perú', 'city': 'Lima', 'year': '2026', 'faculty': 'Ingeniería', 'members': [{'name': 'José García', 'code': 'U123'}]})
    assert response.status_code == 200
    assert response.headers['content-type'] == DOCX_MIME
    doc = Document(BytesIO(response.content))
    assert abs(doc.sections[0].page_width.mm - 210) < .1
    assert abs(doc.sections[0].page_height.mm - 297) < .1
    assert doc.sections[0].left_margin.inches == 1
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Investigación del Perú' in text and 'José García - U123' in text
    assert doc.sections[0].footer.paragraphs[0].text == 'Lima – 2026'
    assert len(doc.inline_shapes) == 1
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text:
                assert run.font.name == 'Calibri' and run.font.size.pt == 11
    cover_id = response.headers['x-cover-id']
    assert client.get(f'/api/covers/{cover_id}/docx').status_code == 200
    assert client.get(f'/api/covers/{cover_id}/pdf').status_code == 200


def test_docx_limits_and_optional_codes():
    response = client.post('/api/docx', json={'show_codes': False, 'members': [{'name': 'María', 'code': 'HIDDEN'}]})
    text = '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert 'María' in text and 'HIDDEN' not in text
    assert client.post('/api/docx', json={'title': 'x' * 351}).status_code == 422
    assert client.get('/api/covers/not-found/docx').status_code == 404
