"""Editable Word cover using ordinary paragraphs and blank lines."""
from io import BytesIO
from functools import lru_cache
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.graphics import renderPM
from reportlab.graphics.shapes import String


def build_docx(drawing, logo):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(drawing.width), Pt(drawing.height)
    section.top_margin = section.bottom_margin = Pt(72)
    section.left_margin = section.right_margin = Pt(72)
    section.footer_distance = Pt(65)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = False
    normal.paragraph_format.keep_with_next = False
    doc.core_properties.title = 'Carátula UTP'
    doc.core_properties.author = ''

    # PNG fallback is supported by Word and other DOCX readers. Text remains editable.
    image = BytesIO(renderPM.drawToString(logo, fmt='PNG', dpi=300))
    logo_height = 275 * logo.height / logo.width
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(logo_height + 3)
    p.add_run().add_picture(image, width=Pt(275), height=Pt(logo_height))
    items = sorted((item for item in drawing.contents if isinstance(item, String)), key=lambda item: -item.y)
    body_items = [item for item in items if item.y > 108]
    footer_items = [item for item in items if item.y <= 108]
    previous_y = None
    for item in body_items:
        if previous_y is not None:
            blank_lines = max(0, round((previous_y - item.y) / 18) - 1)
            for _ in range(blank_lines):
                blank = doc.add_paragraph()
                blank.paragraph_format.line_spacing = 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1
        run = p.add_run(item.text)
        run.font.name, run.font.size = 'Calibri', Pt(11)
        run.bold = item.fontName.endswith('-Bold')
        previous_y = item.y
    if footer_items:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.line_spacing = 1
        for index, item in enumerate(footer_items):
            if index:
                footer.add_run().add_break()
            run = footer.add_run(item.text)
            run.font.name, run.font.size = 'Calibri', Pt(11)
    result = BytesIO()
    doc.save(result)
    return result.getvalue()
